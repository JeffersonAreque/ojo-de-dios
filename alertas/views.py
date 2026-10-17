from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from cuentas.permisos import concesiones_visibles

from .models import Alerta


def alertas_visibles(usuario):
    """El concesionario solo recibe alertas ya confirmadas por el analista (RF-003)."""
    qs = Alerta.objects.filter(concesion__in=concesiones_visibles(usuario)).select_related(
        "concesion", "analisis__imagen_despues__vuelo", "revisado_por"
    )
    if usuario.es_concesionario:
        qs = qs.filter(estado=Alerta.Estado.CONFIRMADA)
    return qs


def alertas_geojson(alertas):
    return [
        {
            "id": a.pk,
            "lat": a.latitud,
            "lon": a.longitud,
            "poligono": a.poligono,
            "gravedad": a.gravedad,
            "estado": a.estado,
            "tipo": a.get_tipo_cambio_display(),
            "area": a.area_m2,
            "fecha": a.fecha_deteccion.strftime("%d/%m/%Y"),
            "concesion": a.concesion.nombre,
        }
        for a in alertas
    ]


@login_required
def historial(request):
    alertas = alertas_visibles(request.user)
    filtros = {k: request.GET.get(k, "") for k in ["estado", "gravedad", "concesion"]}
    if filtros["estado"]:
        alertas = alertas.filter(estado=filtros["estado"])
    if filtros["gravedad"]:
        alertas = alertas.filter(gravedad=filtros["gravedad"])
    if filtros["concesion"]:
        alertas = alertas.filter(concesion_id=filtros["concesion"])

    alertas = alertas.order_by("-analisis__ejecutado", "-area_m2")

    if request.GET.get("formato") == "excel":
        return _excel(alertas)

    return render(request, "alertas/historial.html", {
        "alertas": alertas,
        "filtros": filtros,
        "concesiones": concesiones_visibles(request.user),
        "estados": Alerta.Estado.choices,
        "gravedades": Alerta.Gravedad.choices,
    })


@login_required
def detalle(request, pk):
    alerta = get_object_or_404(alertas_visibles(request.user), pk=pk)

    if request.method == "POST":
        if not request.user.es_analista:
            raise PermissionDenied
        if alerta.estado != Alerta.Estado.PENDIENTE:
            messages.warning(request, "Esta alerta ya fue revisada.")
        else:
            confirmar = request.POST.get("decision") == "confirmar"
            alerta.revisar(request.user, confirmar, request.POST.get("comentario", "").strip())
            messages.success(request, "Alerta confirmada y visible para el concesionario." if confirmar
                             else "Detección descartada; no se notifica al concesionario.")
        siguiente = (
            alertas_visibles(request.user)
            .filter(estado=Alerta.Estado.PENDIENTE, analisis=alerta.analisis)
            .order_by("-area_m2")
            .first()
        )
        return redirect("alertas:detalle", pk=siguiente.pk) if siguiente else redirect("analisis:detalle", pk=alerta.analisis_id)

    c = alerta.concesion
    a = alerta.analisis
    return render(request, "alertas/detalle.html", {
        "alerta": alerta,
        "mapa_concesiones": [{"id": c.pk, "nombre": c.nombre, "limite": c.limite}],
        "mapa_alertas": alertas_geojson([alerta]),
        "mapa_imagenes": [
            {"url": a.imagen_despues.archivo.url, "limites": a.imagen_despues.limites_leaflet, "nombre": "Después"},
            {"url": a.imagen_antes.archivo.url, "limites": a.imagen_antes.limites_leaflet, "nombre": "Antes"},
        ],
    })


def _excel(alertas):
    """Historial de alertas como libro de Excel (.xlsx), igual en cualquier configuracion regional."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Alertas"
    columnas = [
        ("N° alerta", 10), ("Concesión", 26), ("Fecha de detección", 18), ("Tipo de cambio", 28),
        ("Gravedad", 11), ("Estado", 14), ("Área (ha)", 11), ("Latitud", 12), ("Longitud", 12),
        ("Ubicación", 32), ("Revisado por", 16), ("Fecha de revisión", 18), ("Comentario", 60),
    ]
    hoja.append([nombre for nombre, _ in columnas])
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4D2B")
        celda.alignment = Alignment(vertical="center")
    for n, (_, ancho) in enumerate(columnas, start=1):
        hoja.column_dimensions[get_column_letter(n)].width = ancho

    for a in alertas:
        revisor = (a.revisado_por.get_full_name() or a.revisado_por.username) if a.revisado_por else ""
        revision = timezone.localtime(a.fecha_revision).replace(tzinfo=None) if a.fecha_revision else None
        hoja.append([a.pk, a.concesion.nombre, a.fecha_deteccion, a.get_tipo_cambio_display(),
                     a.get_gravedad_display(), a.get_estado_display(), a.area_ha, a.latitud, a.longitud,
                     a.ubicacion_texto, revisor, revision, a.comentario])
        fila = hoja.max_row
        hoja.cell(fila, 3).number_format = "DD/MM/YYYY"
        hoja.cell(fila, 7).number_format = "0.00"
        hoja.cell(fila, 8).number_format = "0.000000"
        hoja.cell(fila, 9).number_format = "0.000000"
        hoja.cell(fila, 12).number_format = "DD/MM/YYYY HH:MM"
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    respuesta = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    respuesta["Content-Disposition"] = f'attachment; filename="historial_alertas_{timezone.localdate():%Y%m%d}.xlsx"'
    libro.save(respuesta)
    return respuesta
